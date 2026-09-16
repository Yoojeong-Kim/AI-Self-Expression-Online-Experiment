import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

csv_path = r'C:\Users\COM\.gemini\antigravity\brain\bb7f545d-8c6c-4d95-bbf7-7ca05adfdffa\survey_codebook.csv'
excel_path = r'C:\Users\COM\Documents\성균관\AXIS\자유 연구\AI Selfie 연구\AI Selfie_code book.xlsx'

df = pd.read_csv(csv_path)

with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    df.to_excel(writer, index=False, sheet_name='Codebook')
    
    workbook = writer.book
    worksheet = writer.sheets['Codebook']
    
    # Header formatting
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
    # Content formatting & Column width
    worksheet.column_dimensions['A'].width = 30
    worksheet.column_dimensions['B'].width = 120
    
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center')

print('Success')
